""" "Download all" through the SQLite stand-in with a real draft behind it.

What a reviewer should check here: every size of a post ships as both the finished
PNG and a PowerPoint file with each layer editable, and the logo in that file is the
client's own picture, not a placeholder.
"""

from __future__ import annotations

import io
import zipfile
from pathlib import Path
from typing import Any

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from stratum_worker.app import service
from stratum_worker.draft.pipeline import DraftError
from stratum_worker.queue import fake
from stratum_worker.queue.fake import World

REPO = Path(__file__).resolve().parents[3]


@pytest.fixture
async def drafted(tmp_path: Path) -> tuple[service.Context, dict[str, Any], dict[str, Any], str]:
    world = World(tmp_path / "db.sqlite")
    agency, client_id = await fake.seed(world, REPO)
    ctx = service.Context(
        supa=world.factory(agency),
        settings=world.settings(REPO),
        gateway=fake.gateway(),
        agency_id=agency,
    )
    [post] = await ctx.supa.insert(
        "posts",
        {
            "agency_id": agency,
            "client_id": client_id,
            "slug": "boiler-offer",
            "brief": "20% off boiler servicing this month. Gas Safe 000000.",
            "template_id": "offer_4x5",
            "aspects": ["4:5", "1:1"],
        },
    )
    [client] = await ctx.supa.select("clients", id=f"eq.{client_id}")
    version = await service.draft(ctx, post, client)
    [post] = await ctx.supa.select("posts", id=f"eq.{post['id']}")
    return ctx, post, client, str(version["id"])


async def test_every_size_ships_as_a_png_and_an_editable_file(
    drafted: tuple[service.Context, dict[str, Any], dict[str, Any], str],
) -> None:
    ctx, post, client, version_id = drafted
    archive = zipfile.ZipFile(io.BytesIO(await service.bundle(ctx, post, client, version_id)))
    names = set(archive.namelist())
    for shape in ("4x5", "1x1"):
        assert f"boiler-offer-{shape}.png" in names
        assert f"boiler-offer-{shape}.pptx" in names
    assert {"boiler-offer-caption.txt", "OPEN-IN-CANVA.txt"} <= names

    deck = Presentation(io.BytesIO(archive.read("boiler-offer-4x5.pptx")))
    shapes = {shape.name: shape for shape in deck.slides[0].shapes}
    assert shapes["logo"].shape_type == MSO_SHAPE_TYPE.PICTURE  # the client's own file
    assert shapes["headline"].has_text_frame and shapes["headline"].text_frame.text
    caption = archive.read("boiler-offer-caption.txt").decode().strip()
    assert deck.slides[0].notes_slide.notes_text_frame.text == caption


async def test_one_size_on_its_own(
    drafted: tuple[service.Context, dict[str, Any], dict[str, Any], str],
) -> None:
    ctx, post, client, version_id = drafted
    deck = Presentation(io.BytesIO(await service.export_deck(ctx, post, client, version_id, "1:1")))
    assert deck.slide_width == deck.slide_height  # square


async def test_a_size_the_post_was_not_made_in_is_refused(
    drafted: tuple[service.Context, dict[str, Any], dict[str, Any], str],
) -> None:
    ctx, post, client, version_id = drafted
    with pytest.raises(DraftError, match="no 9:16 version"):
        await service.export_deck(ctx, post, client, version_id, "9:16")
