"""The PowerPoint export: every layer an editable object, drawn where the poster has it.

What a reviewer should check here:
- one PowerPoint object per visible layer, in the poster's stacking order, named
  after the layer, so the designer in Canva recognises each one;
- text lands at the renderer's own size and line breaks (the same layout code);
- a photo is cropped exactly as the poster crops it; a logo is never cropped;
- brand colours and opacities carry over, with no theme styling added on top.
"""

from __future__ import annotations

import hashlib
import io
from itertools import pairwise
from pathlib import Path
from typing import Any

import pytest
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from stratum_spec import (
    BrandKit,
    ImageLayer,
    LogoLayer,
    PosterSpec,
    ShapeLayer,
    TextLayer,
)
from stratum_worker.export import export_pptx
from stratum_worker.export.pptx import EMU_PER_PX, MAX_SCALE, PT_PER_PX
from stratum_worker.render.layout import place_text
from stratum_worker.render.svg import PLACEHOLDER_FILL

TEXTS = {
    "headline": "20% off boiler servicing",
    "offer": "Book before 31 October. Gas Safe registered engineers.",
    "cta": "Call 0161 496 0000",
}


def _sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


@pytest.fixture
def photo(tmp_path: Path) -> Path:
    image = Image.new("RGB", (1600, 1200), (40, 90, 140))
    ImageDraw.Draw(image).rectangle([600, 300, 1000, 1000], fill=(235, 235, 230))
    path = tmp_path / "photo.png"
    image.save(path)
    return path


@pytest.fixture
def logo(repo_root: Path) -> Path:
    return repo_root / "clients" / "kowalski-plumbing" / "logo.png"


def _filled(spec: PosterSpec, **assets: str) -> PosterSpec:
    """The template with real words and the given assets bound, by layer type."""
    layers: list[dict[str, Any]] = []
    for layer in spec.layers:
        data = layer.model_dump()
        if isinstance(layer, TextLayer):
            data["content"] = TEXTS.get(layer.id, TEXTS["headline"])
        elif isinstance(layer, ImageLayer) and "photo" in assets:
            data["asset"] = assets["photo"]
        elif isinstance(layer, LogoLayer) and "logo" in assets:
            data["asset"] = assets["logo"]
        layers.append(data)
    return PosterSpec.model_validate({**spec.model_dump(), "layers": layers})


def _deck(spec: PosterSpec, kit: BrandKit, root: Path, **kwargs: Any) -> Any:
    return Presentation(io.BytesIO(export_pptx(spec, kit, root=root, **kwargs)))


def _shapes(deck: Any) -> dict[str, Any]:
    return {shape.name: shape for shape in deck.slides[0].shapes}


@pytest.fixture
def poster(offer: PosterSpec, photo: Path, logo: Path) -> tuple[PosterSpec, dict[str, str]]:
    assets = {_sha(photo): str(photo), _sha(logo): str(logo)}
    return _filled(offer, photo=_sha(photo), logo=_sha(logo)), assets


def test_the_slide_is_exactly_the_poster(
    poster: tuple[PosterSpec, dict[str, str]], kit: BrandKit, repo_root: Path
) -> None:
    spec, assets = poster
    deck = _deck(spec, kit, repo_root, assets=assets)
    width, height = spec.canvas.size
    assert (deck.slide_width, deck.slide_height) == (width * EMU_PER_PX, height * EMU_PER_PX)
    assert len(deck.slides) == 1


def test_one_object_per_visible_layer_in_the_posters_order(
    poster: tuple[PosterSpec, dict[str, str]], kit: BrandKit, repo_root: Path
) -> None:
    spec, assets = poster
    expected: list[str] = []
    for layer in spec.layers:
        expected.append(layer.id)
        if isinstance(layer, ImageLayer) and layer.dim > 0:
            expected.append(f"{layer.id} dim")  # removable, above its photo
    deck = _deck(spec, kit, repo_root, assets=assets)
    assert [shape.name for shape in deck.slides[0].shapes] == expected


def test_text_is_the_renderers_own_size_and_line_breaks(
    poster: tuple[PosterSpec, dict[str, str]], kit: BrandKit, repo_root: Path
) -> None:
    spec, assets = poster
    shapes = _shapes(_deck(spec, kit, repo_root, assets=assets))
    for layer in spec.layers:
        if not isinstance(layer, TextLayer):
            continue
        placed = place_text(layer, spec, kit, repo_root)
        frame = shapes[layer.id].text_frame
        assert frame.text == "\v".join(placed.lines)  # "\v" is a line break in one paragraph
        [paragraph] = frame.paragraphs
        for run in paragraph.runs:
            assert run.font.name == placed.family
            assert run.font.size.pt == pytest.approx(placed.size_px * PT_PER_PX, abs=0.01)
            assert run.font.bold == (placed.weight >= 600)
            assert f"#{run.font.color.rgb}".lower() == placed.color.lower()
        assert paragraph.line_spacing.pt == pytest.approx(
            placed.line_height_px * PT_PER_PX, abs=0.01
        )


def test_a_photo_is_cropped_exactly_as_the_poster_crops_it(
    poster: tuple[PosterSpec, dict[str, str]], kit: BrandKit, repo_root: Path
) -> None:
    spec, assets = poster
    background = _shapes(_deck(spec, kit, repo_root, assets=assets))["bg"]
    # A 4:3 photo in a 4:5 box: the poster scales it to the box's height and trims
    # the sides evenly, keeping 0.8 / (4/3) = 60% of the width.
    assert background.shape_type == MSO_SHAPE_TYPE.PICTURE
    assert background.crop_left == pytest.approx(0.2, abs=1e-4)
    assert background.crop_right == pytest.approx(0.2, abs=1e-4)
    assert background.crop_top == background.crop_bottom == 0


def test_the_logo_is_never_cropped_and_keeps_its_proportions(
    poster: tuple[PosterSpec, dict[str, str]], kit: BrandKit, repo_root: Path, logo: Path
) -> None:
    spec, assets = poster
    picture = _shapes(_deck(spec, kit, repo_root, assets=assets))["logo"]
    assert (picture.crop_left, picture.crop_right, picture.crop_top, picture.crop_bottom) == (
        0,
        0,
        0,
        0,
    )
    with Image.open(logo) as original:
        ratio = original.width / original.height
    assert picture.width / picture.height == pytest.approx(ratio, rel=0.01)
    # Centred in its box, like the renderer's "meet".
    box = spec.layer("logo").box
    width, height = spec.canvas.size
    centre_x = (box.x + box.w / 2) * width * EMU_PER_PX
    centre_y = (box.y + box.h / 2) * height * EMU_PER_PX
    assert picture.left + picture.width / 2 == pytest.approx(centre_x, abs=EMU_PER_PX)
    assert picture.top + picture.height / 2 == pytest.approx(centre_y, abs=EMU_PER_PX)


def test_a_transparent_logo_stays_transparent(
    poster: tuple[PosterSpec, dict[str, str]], kit: BrandKit, repo_root: Path, logo: Path
) -> None:
    spec, assets = poster
    picture = _shapes(_deck(spec, kit, repo_root, assets=assets))["logo"]
    with Image.open(logo) as original:
        transparent = original.mode in ("RGBA", "LA", "P")
    embedded = Image.open(io.BytesIO(picture.image.blob))
    assert (embedded.format == "PNG") == transparent


def test_shapes_carry_the_brand_colour_and_opacity_and_nothing_else(
    poster: tuple[PosterSpec, dict[str, str]], kit: BrandKit, repo_root: Path
) -> None:
    spec, assets = poster
    shapes = _shapes(_deck(spec, kit, repo_root, assets=assets))
    for layer in spec.layers:
        if not isinstance(layer, ShapeLayer):
            continue
        shape = shapes[layer.id]
        assert f"#{shape.fill.fore_color.rgb}".lower() == kit.color(layer.fill).lower()
        alpha = shape._element.spPr.find(f".//{qn('a:alpha')}")
        if layer.opacity < 1:
            assert alpha is not None and int(alpha.get("val")) == round(layer.opacity * 100000)
        else:
            assert alpha is None
        # No theme styling: an importer must not add a shadow or an outline.
        assert shape._element.find(qn("p:style")) is None


def test_hidden_layers_are_left_out(
    poster: tuple[PosterSpec, dict[str, str]], kit: BrandKit, repo_root: Path
) -> None:
    spec, assets = poster
    hidden = PosterSpec.model_validate(
        {
            **spec.model_dump(),
            "layers": [
                {**layer.model_dump(), "visible": layer.id != "cta"} for layer in spec.layers
            ],
        }
    )
    assert "cta" not in _shapes(_deck(hidden, kit, repo_root, assets=assets))


def test_an_empty_slot_is_a_grey_box_that_says_so(
    offer: PosterSpec, kit: BrandKit, repo_root: Path
) -> None:
    shapes = _shapes(_deck(_filled(offer), kit, repo_root))
    placeholder = shapes["bg (empty image)"]
    assert f"#{placeholder.fill.fore_color.rgb}".lower() == PLACEHOLDER_FILL
    assert "logo (empty logo)" in shapes


def test_native_shapes_rotation_and_rounded_corners(kit: BrandKit, repo_root: Path) -> None:
    spec = PosterSpec.model_validate(
        {
            "template_id": "shapes",
            "layers": [
                {
                    "type": "shape",
                    "id": "chip",
                    "shape": "rounded",
                    "radius": 0.2,
                    "box": {"x": 0.1, "y": 0.1, "w": 0.4, "h": 0.2, "rotation": 15},
                },
                {
                    "type": "shape",
                    "id": "dot",
                    "shape": "ellipse",
                    "box": {"x": 0.6, "y": 0.6, "w": 0.2, "h": 0.2},
                },
            ],
        }
    )
    shapes = _shapes(_deck(spec, kit, repo_root))
    assert shapes["chip"].auto_shape_type == MSO_SHAPE.ROUNDED_RECTANGLE
    assert shapes["chip"].adjustments[0] == pytest.approx(0.2)
    assert shapes["chip"].rotation == pytest.approx(15)
    assert shapes["dot"].auto_shape_type == MSO_SHAPE.OVAL


def test_a_blurred_photo_arrives_blurred(
    offer: PosterSpec, kit: BrandKit, repo_root: Path, photo: Path
) -> None:
    sharp = _filled(offer, photo=_sha(photo))
    blurred = PosterSpec.model_validate(
        {
            **sharp.model_dump(),
            "layers": [
                {**layer.model_dump(), **({"blur": 0.5} if layer.id == "bg" else {})}
                for layer in sharp.layers
            ],
        }
    )
    assets = {_sha(photo): str(photo)}

    def steepest(spec: PosterSpec) -> int:
        picture = _shapes(_deck(spec, kit, repo_root, assets=assets))["bg"]
        pixels = Image.open(io.BytesIO(picture.image.blob)).convert("L").tobytes()
        # The photo has one hard-edged box. Blur spreads that edge over many pixels,
        # so the steepest step between neighbours drops sharply.
        return max(abs(a - b) for a, b in pairwise(pixels))

    assert steepest(blurred) < steepest(sharp) / 4


def test_a_huge_photo_is_stored_at_most_twice_its_shown_size(
    offer: PosterSpec, kit: BrandKit, repo_root: Path, tmp_path: Path
) -> None:
    big = tmp_path / "big.png"
    Image.new("RGB", (6000, 4500), (10, 20, 30)).save(big)
    spec = _filled(offer, photo=_sha(big))
    picture = _shapes(_deck(spec, kit, repo_root, assets={_sha(big): str(big)}))["bg"]
    embedded = Image.open(io.BytesIO(picture.image.blob))
    _, height = spec.canvas.size
    # Cover-scaled to the box's height (1350 px), so kept at up to 2x that.
    assert embedded.height <= height * MAX_SCALE + 1
    assert embedded.format == "JPEG"  # no transparency to keep: a tenth of the size


def test_the_caption_travels_in_the_notes(
    poster: tuple[PosterSpec, dict[str, str]], kit: BrandKit, repo_root: Path
) -> None:
    spec, assets = poster
    deck = _deck(spec, kit, repo_root, assets=assets, notes="Book before winter.\n\n#boiler")
    assert deck.slides[0].notes_slide.notes_text_frame.text == "Book before winter.\n\n#boiler"
