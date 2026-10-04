"""Words behind the subject (M10g, spec schema 6).

A head rising into the headline looks like a mistake when the letters are printed
over it, and intended - a magazine cover - when the head is in front of them. The
headline goes behind the subject only where a head already reaches it and every
word can still be read; the renderer and the PowerPoint export then draw the
subject again over the headline.
"""

from __future__ import annotations

import io
import json
from pathlib import Path

import numpy as np
import pytest
from PIL import Image, ImageDraw, ImageFilter
from pptx import Presentation
from stratum_spec import (
    Box,
    BrandKit,
    Canvas,
    Crop,
    ImageLayer,
    PosterSpec,
    ShapeLayer,
    TextLayer,
)
from stratum_worker.app import service
from stratum_worker.design.choose import Material, fill
from stratum_worker.design.families import COVER_PHOTO
from stratum_worker.export import export_pptx
from stratum_worker.photo import depth
from stratum_worker.photo.depth import decide_depth, hidden, letters
from stratum_worker.photo.faces import Face
from stratum_worker.photo.frame import face_flags, frame_photos, framing_cost
from stratum_worker.photo.read import PhotoRead
from stratum_worker.photo.subject import guided, matte, subject_of
from stratum_worker.queue import fake
from stratum_worker.queue.fake import World
from stratum_worker.render import render_poster
from stratum_worker.render.layout import box_px

REPO = Path(__file__).resolve().parents[3]
KIT = BrandKit.model_validate(
    json.loads((REPO / "clients/kowalski-plumbing/brand.json").read_text())
)
PHOTO, MATTE, LOGO = "a" * 64, "d" * 64, "c" * 64
W, H = 1000, 1250  # the poster's own shape (4:5): no crop moves anything


# --- a person, drawn -------------------------------------------------------------------


def person(head_y: int = 330, radius: int = 130) -> Image.Image:
    """The subject's outline: a head and the shoulders below it, white on black."""
    outline = Image.new("L", (W, H), 0)
    draw = ImageDraw.Draw(outline)
    draw.ellipse([500 - radius, head_y - radius, 500 + radius, head_y + radius], 255)
    draw.rectangle([330, head_y + radius - 20, 670, H], 255)
    return outline


def photo_of(outline: Image.Image) -> Image.Image:
    """A grey wall, the person in red where the outline is."""
    wall = Image.new("RGB", (W, H), (200, 200, 200))
    wall.paste(Image.new("RGB", (W, H), (200, 30, 30)), mask=outline)
    return wall


def read_of(
    outline: Image.Image, *, head_y: int = 330, radius: int = 130, faces: bool = True
) -> PhotoRead:
    small = np.asarray(outline.resize((128, 160)), dtype=np.float32) / 255.0
    face = Face(
        x=(500 - 0.75 * radius) / W,
        y=(head_y - 0.6 * radius) / H,
        w=1.5 * radius / W,
        h=1.2 * radius / H,
        score=0.95,
    )
    return PhotoRead(
        size=(W, H),
        faces=(face,) if faces else (),
        salience=np.full((40, 32), 1 / 1280),
        busy=np.zeros((40, 32)),
        strength=0.0,
        subject=subject_of(small),
    )


def poster(*, foot_first: bool = False) -> PosterSpec:
    """A photo under a band with the headline on it; the offer on a band at the foot."""
    photo = ImageLayer(id="bg", box=Box(x=0, y=0, w=1, h=1), asset=PHOTO)
    band = ShapeLayer(
        id="band", box=Box(x=0, y=0, w=1, h=0.26), fill="palette.primary", opacity=0.92
    )
    headline = TextLayer(
        id="headline",
        box=Box(x=0.06, y=0.04, w=0.88, h=0.2),
        content="Winter boiler checks",
        font="font.display",
        color="palette.text_on_primary",
        align="center",
        uppercase=True,
        max_size=0.12,
        min_size=0.05,
        max_lines=2,
    )
    foot = ShapeLayer(
        id="foot", box=Box(x=0, y=0.84, w=1, h=0.16), fill="palette.primary", opacity=0.92
    )
    offer = TextLayer(
        id="offer",
        box=Box(x=0.06, y=0.86, w=0.88, h=0.06),
        content="20% off until 31 October",
        font="font.body",
        color="palette.text_on_primary",
        align="center",
        max_size=0.04,
    )
    layers = (
        [photo, foot, band, headline, offer] if foot_first else [photo, band, headline, foot, offer]
    )
    return PosterSpec(template_id="manual", canvas=Canvas(aspect="4:5"), layers=layers)


def behind_of(spec: PosterSpec) -> str | None:
    headline = spec.layer("headline")
    assert isinstance(headline, TextLayer)
    return headline.behind


def decided(spec: PosterSpec, read: PhotoRead) -> PosterSpec:
    return decide_depth(spec, {PHOTO: read}, KIT, REPO, ["4:5"])


# --- when the words go behind ----------------------------------------------------------


def test_a_head_rising_into_the_headline_goes_in_front_of_it() -> None:
    assert behind_of(decided(poster(), read_of(person()))) == "bg"


@pytest.mark.parametrize(
    ("head_y", "radius", "why"),
    [
        (480, 130, "the head does not reach the letters"),
        (260, 130, "it would hide more than a quarter of the letters"),
        (230, 150, "it would hide nearly half of them"),
    ],
)
def test_the_words_stay_in_front_otherwise(head_y: int, radius: int, why: str) -> None:
    read = read_of(person(head_y, radius), head_y=head_y, radius=radius)
    assert behind_of(decided(poster(), read)) is None, why


def test_only_a_person_goes_in_front_not_a_product() -> None:
    assert behind_of(decided(poster(), read_of(person(), faces=False))) is None


def test_not_when_the_photo_has_no_clear_subject() -> None:
    read = read_of(person())
    vague = PhotoRead(
        size=read.size, faces=read.faces, salience=read.salience, busy=read.busy, strength=0.0
    )
    assert decided(poster(), vague) == poster()


def test_not_when_it_would_take_away_the_panel_other_words_are_read_on() -> None:
    # The same offer, but its band is drawn before the headline: the person drawn
    # again would cover the band, and the offer would sit on their jacket.
    assert behind_of(decided(poster(foot_first=True), read_of(person()))) is None


def test_not_when_one_line_would_lose_too_much(monkeypatch: pytest.MonkeyPatch) -> None:
    lines = [(100.0, 100.0, 800.0, 90.0)] * 3
    monkeypatch.setattr(depth, "letters", lambda *args: lines)
    monkeypatch.setattr(depth, "_head_meets", lambda *args: True)
    monkeypatch.setattr(depth, "hidden", lambda *args: [0.0, 0.0, 0.45])
    # 15% of the letters overall, but 45% of the last line.
    assert behind_of(decided(poster(), read_of(person()))) is None


def test_it_must_fit_in_every_size(monkeypatch: pytest.MonkeyPatch) -> None:
    spec = poster().model_copy(update={"variants": {"1:1": {}}})
    verdicts = iter([{"bg": "fits"}, {"bg": "refused"}])
    monkeypatch.setattr(depth, "_verdict", lambda *args: next(verdicts))
    chosen = decide_depth(spec, {PHOTO: read_of(person())}, KIT, REPO, ["4:5", "1:1"])
    assert behind_of(chosen) is None


def test_it_is_decided_again_from_scratch() -> None:
    spec = decided(poster(), read_of(person()))
    assert behind_of(spec) == "bg"
    far = read_of(person(480), head_y=480)
    assert behind_of(decided(spec, far)) is None


def test_framing_decides_it_after_choosing_the_crop() -> None:
    framed = frame_photos(poster(), {PHOTO: read_of(person())}, KIT, REPO, ["4:5"])
    assert behind_of(framed) == "bg"
    photo = framed.layer("bg")
    assert isinstance(photo, ImageLayer) and photo.crop is not None


# --- the cover family: where depth happens --------------------------------------------


def material(read: PhotoRead) -> Material:
    return Material(
        words={
            "headline": "Winter boiler checks",
            "offer_detail": "20% off until 31 October.",
            "cta": "Book today",
        },
        kit=KIT,
        photos=(PHOTO,),
        background=None,
        logo=LOGO,
        logo_aspect=2.0,
        ink=None,
        reads={PHOTO: read},
    )


def test_a_cover_has_the_headlines_band_below_it_and_the_details_band_above() -> None:
    spec = fill("ds.cover.1.0001", material(read_of(person())), ["4:5"], REPO)
    assert spec is not None
    order = [layer.id for layer in spec.layers]
    assert order.index("bg") < order.index("band") < order.index("headline")
    assert order.index("headline") < order.index("backing") < order.index("offer")
    _, height = spec.canvas.size
    bands = sum(box_px(spec, spec.layer(one))[3] for one in ("band", "backing"))
    assert bands <= (1 - COVER_PHOTO) * height + 1


def test_a_cover_puts_the_headline_behind_a_head_that_reaches_it() -> None:
    stuff = material(read_of(person(200, 110), head_y=200, radius=110))
    calm = fill("ds.cover.1.0001", stuff, ["4:5"], REPO)
    bold = fill("ds.cover.1.0000", stuff, ["4:5"], REPO)
    assert calm is not None and bold is not None
    assert behind_of(frame_photos(calm, stuff.reads, KIT, REPO, ["4:5"])) == "bg"
    # The bold headline is bigger and lower: the head would hide too much of it.
    assert behind_of(frame_photos(bold, stuff.reads, KIT, REPO, ["4:5"])) is None


# --- what it measures ------------------------------------------------------------------


def test_letters_are_measured_line_by_line_inside_the_box() -> None:
    spec = poster()
    headline = spec.layer("headline")
    assert isinstance(headline, TextLayer)
    lines = letters(headline, spec, KIT, REPO)
    assert len(lines) == 2
    width, height = spec.canvas.size
    for x, y, w, h in lines:
        assert 0.06 * width - 1 <= x and x + w <= 0.94 * width + 1
        assert 0.04 * height - 1 <= y and y + h <= 0.24 * height + 1


def test_how_much_of_the_letters_the_subject_covers_follows_the_crop() -> None:
    spec = poster()
    photo = spec.layer("bg")
    headline = spec.layer("headline")
    assert isinstance(photo, ImageLayer) and isinstance(headline, TextLayer)
    lines = letters(headline, spec, KIT, REPO)
    read = read_of(person())
    assert hidden(spec, photo, read, lines)[0] == 0.0
    assert 0.1 < hidden(spec, photo, read, lines)[1] < 0.4
    # Zoomed in on the lower part of the photo, the head is larger and higher in the
    # poster: more of the letters are covered.
    zoomed = photo.model_copy(update={"crop": Crop(x=0.5, y=1.0, zoom=1.3)})
    assert hidden(spec, zoomed, read, lines)[1] > hidden(spec, photo, read, lines)[1]


def test_with_depth_the_head_no_longer_counts_as_covered() -> None:
    reads = {PHOTO: read_of(person())}
    plain = poster()
    deep = decided(plain, reads[PHOTO])
    assert framing_cost(deep, reads, KIT, REPO) < framing_cost(plain, reads, KIT, REPO)
    assert face_flags(deep, reads, KIT, REPO) == ()


# --- drawing it ------------------------------------------------------------------------


@pytest.fixture
def files(tmp_path: Path) -> dict[str, str]:
    outline = person()
    photo, matted = tmp_path / "photo.png", tmp_path / "matte.png"
    photo_of(outline).save(photo)
    outline.save(matted)
    return {PHOTO: str(photo), MATTE: str(matted)}


def with_matte(spec: PosterSpec) -> PosterSpec:
    layers = [
        layer.model_copy(update={"subject": MATTE}) if isinstance(layer, ImageLayer) else layer
        for layer in spec.layers
    ]
    return spec.model_copy(update={"layers": layers})


def pixels(spec: PosterSpec, files: dict[str, str]) -> np.ndarray:  # type: ignore[type-arg]
    png = render_poster(spec, KIT, assets=files, root=REPO)
    return np.asarray(Image.open(io.BytesIO(png)).convert("RGB"), dtype=np.int16)


def test_the_subject_is_drawn_over_the_headline_and_its_band(files: dict[str, str]) -> None:
    deep = with_matte(decided(poster(), read_of(person())))
    shown = pixels(deep, files)
    front = pixels(poster(), files)
    bare = pixels(poster().model_copy(update={"layers": poster().layers[:1]}), files)
    big = person().resize((1080, 1350))
    # A few pixels in from the outline's edge either way: edges are anti-aliased.
    subject = np.asarray(big.filter(ImageFilter.MinFilter(9))) >= 250
    background = np.asarray(big.filter(ImageFilter.MaxFilter(9))) <= 5
    top = np.zeros_like(subject)
    top[: int(0.3 * 1350)] = True  # the band and the headline
    # On the subject, the band and the letters are gone: the photo itself shows.
    assert np.abs(shown - bare)[subject & top].max() <= 2
    # Everywhere else, exactly the poster with the words in front.
    assert np.abs(shown - front)[background].max() <= 2
    # And it really changed something: letters were over the head.
    assert np.abs(front - bare)[subject & top].max() > 50


def test_without_its_matte_the_words_are_drawn_in_front(files: dict[str, str]) -> None:
    deep = decided(poster(), read_of(person()))  # behind, but no matte yet
    assert np.array_equal(pixels(deep, files), pixels(poster(), files))


def test_a_hidden_photo_brings_nothing_in_front(files: dict[str, str]) -> None:
    deep = with_matte(decided(poster(), read_of(person())))
    hidden_photo = deep.model_copy(
        update={
            "layers": [
                layer.model_copy(update={"visible": False}) if layer.id == "bg" else layer
                for layer in deep.layers
            ]
        }
    )
    plain = poster().model_copy(
        update={
            "layers": [
                layer.model_copy(update={"visible": False}) if layer.id == "bg" else layer
                for layer in poster().layers
            ]
        }
    )
    assert np.array_equal(pixels(hidden_photo, files), pixels(plain, files))


def test_the_powerpoint_puts_the_subject_in_front_of_the_headline(files: dict[str, str]) -> None:
    deep = with_matte(decided(poster(), read_of(person())))
    deck = Presentation(io.BytesIO(export_pptx(deep, KIT, assets=files, root=REPO)))
    names = [shape.name for shape in deck.slides[0].shapes]
    front = "headline behind: subject of bg"
    assert names.index("headline") < names.index(front) < names.index("foot")
    photo = next(s for s in deck.slides[0].shapes if s.name == "bg")
    cut = next(s for s in deck.slides[0].shapes if s.name == front)
    assert (cut.left, cut.top, cut.width, cut.height) == (
        photo.left,
        photo.top,
        photo.width,
        photo.height,
    )
    assert cut.crop_top == pytest.approx(photo.crop_top)
    assert cut.crop_left == pytest.approx(photo.crop_left)
    plain = Presentation(io.BytesIO(export_pptx(poster(), KIT, assets=files, root=REPO)))
    assert front not in [shape.name for shape in plain.slides[0].shapes]


# --- the matte -------------------------------------------------------------------------


def test_the_guided_filter_snaps_a_soft_edge_to_the_photos_edge() -> None:
    guide = np.zeros((60, 60))
    guide[:, 30:] = 1.0  # a hard edge in the photo
    soft = np.clip((np.arange(60) - 20) / 20, 0, 1)[None, :].repeat(60, 0)  # a blurry mask
    out = guided(guide, soft, 3, 1e-3)
    # Across the photo's edge the mask now jumps, where it used to creep.
    assert (out[:, 30] - out[:, 29]).mean() > 3 * (soft[:, 30] - soft[:, 29]).mean()
    assert np.allclose(out[:, :15], 0.0, atol=0.01) and np.allclose(out[:, 45:], 1.0, atol=0.01)
    flat = guided(guide, np.full((60, 60), 0.4), 3, 1e-3)
    assert np.allclose(flat, 0.4, atol=1e-6)


def test_the_matte_has_the_photos_shape_and_firm_edges() -> None:
    image = Image.new("RGB", (900, 600), (195, 195, 195))
    ImageDraw.Draw(image).ellipse([150, 200, 350, 400], (200, 30, 30))
    made = matte(image)
    assert made.mode == "L" and made.size == (900, 600)
    values = np.asarray(made)
    assert values[300, 250] > 240  # the disc
    assert values[50, 800] < 15  # the wall
    assert (values > 15).sum() - (values > 240).sum() < 0.01 * values.size  # few in-betweens


def test_a_large_photo_gets_a_matte_no_larger_than_1600() -> None:
    made = matte(Image.new("RGB", (3200, 2000), (120, 120, 120)))
    assert made.size == (1600, 1000)


# --- through the service ---------------------------------------------------------------


@pytest.fixture
async def ctx(tmp_path: Path) -> service.Context:
    world = World(tmp_path / "db.sqlite")
    agency, _ = await fake.seed(world, REPO)
    return service.Context(
        supa=world.factory(agency),
        settings=world.settings(REPO),
        gateway=fake.gateway(),
        agency_id=agency,
    )


async def test_a_poster_with_depth_gets_its_matte_stored_once(
    ctx: service.Context, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(service, "_MATTES", {})
    sha = await service.store_image(ctx, None, _png(photo_of(person())), kind="photo")
    spec = decided(poster(), read_of(person()))
    spec = spec.model_copy(
        update={
            "layers": [
                layer.model_copy(update={"asset": sha}) if layer.id == "bg" else layer
                for layer in spec.layers
            ]
        }
    )
    paths: dict[str, str] = {}
    first = await service.with_mattes(ctx, spec, paths)
    photo = first.layer("bg")
    assert isinstance(photo, ImageLayer) and photo.subject is not None
    assert photo.subject in paths and Path(paths[photo.subject]).is_file()
    # Kept in storage, so another machine can draw it.
    stored = await ctx.supa.download(f"{ctx.agency_id}/assets/{photo.subject}.png")
    with Image.open(io.BytesIO(stored)) as image:
        assert image.size == (W, H)
    # Made once: the same photo again costs nothing.
    calls: list[str] = []
    monkeypatch.setattr(service, "_matte_png", lambda path: calls.append(path) or b"")
    again = await service.with_mattes(ctx, spec, {})
    assert again.layer("bg") == photo and not calls


async def test_a_poster_without_depth_is_left_alone(ctx: service.Context) -> None:
    spec = poster()
    assert await service.with_mattes(ctx, spec, {}) is spec


def _png(image: Image.Image) -> bytes:
    out = io.BytesIO()
    image.save(out, "PNG")
    return out.getvalue()
