"""The logo's panel, as drawn by the renderer and as exported to Canva.

The promises:
- the panel hugs the logo as drawn, and logo and panel stay inside the logo's box,
  so a panel never runs into the text beside it;
- the logo itself is drawn unchanged, just smaller to make room;
- the PowerPoint file has the panel as its own rounded shape under the logo, so a
  designer in Canva can restyle or delete it.
"""

from __future__ import annotations

import io
import json
from pathlib import Path

from logo_samples import draw, two_colour
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from stratum_spec import BrandKit, LogoLayer, LogoPlate, PosterSpec
from stratum_spec.models import Box, Canvas, ShapeLayer
from stratum_worker.export import export_pptx
from stratum_worker.render import render_poster
from stratum_worker.render.layout import box_px, plated

REPO = Path(__file__).resolve().parents[3]
KIT = BrandKit.model_validate(
    json.loads((REPO / "clients" / "kowalski-plumbing" / "brand.json").read_text())
)
SHA = "b" * 64
PLATE = LogoPlate(fill="palette.surface", opacity=1.0)


def spec_with(plate: LogoPlate | None) -> PosterSpec:
    return PosterSpec(
        template_id="test",
        canvas=Canvas(aspect="1:1", base_width=600),
        layers=[
            ShapeLayer(id="band", box=Box(x=0, y=0, w=1, h=1), fill="palette.primary"),
            LogoLayer(id="logo", box=Box(x=0.1, y=0.1, w=0.8, h=0.4), asset=SHA, plate=plate),
        ],
    )


def logo_file(tmp_path: Path) -> tuple[dict[str, str], tuple[int, int]]:
    path = tmp_path / "logo.png"
    path.write_bytes(draw(two_colour()))
    with Image.open(path) as image:
        return {SHA: str(path)}, image.size


def test_the_panel_hugs_the_logo_and_stays_inside_its_box(tmp_path: Path) -> None:
    assets, size = logo_file(tmp_path)
    spec = spec_with(PLATE)
    box = box_px(spec, spec.layer("logo"))
    logo, panel, radius = plated(box, size, PLATE)
    x, y, w, h = box
    px, py, pw, ph = panel
    assert x - 1e-6 <= px and y - 1e-6 <= py
    assert px + pw <= x + w + 1e-6 and py + ph <= y + h + 1e-6
    # Same shape as the file: the logo is scaled, never stretched.
    assert abs(logo[2] / logo[3] - size[0] / size[1]) < 1e-6
    assert radius > 0

    png = Image.open(io.BytesIO(render_poster(spec, KIT, assets=assets, root=REPO))).convert("RGB")
    surface = tuple(int(KIT.palette["surface"][i : i + 2], 16) for i in (1, 3, 5))
    primary = tuple(int(KIT.palette["primary"][i : i + 2], 16) for i in (1, 3, 5))
    # Inside the panel's padding (between its edge and the logo): the page colour.
    assert png.getpixel((int(px + pw / 2), int(py + 2))) == surface
    # Just outside the panel: still the band.
    assert png.getpixel((int(px + pw / 2), max(0, int(py) - 3))) == primary


def test_without_a_panel_nothing_changes(tmp_path: Path) -> None:
    assets, _ = logo_file(tmp_path)
    plain = render_poster(spec_with(None), KIT, assets=assets, root=REPO)
    assert render_poster(spec_with(None), KIT, assets=assets, root=REPO) == plain
    assert render_poster(spec_with(PLATE), KIT, assets=assets, root=REPO) != plain


def test_the_canva_file_has_the_panel_as_its_own_shape_under_the_logo(tmp_path: Path) -> None:
    assets, _ = logo_file(tmp_path)
    deck = Presentation(io.BytesIO(export_pptx(spec_with(PLATE), KIT, assets=assets, root=REPO)))
    shapes = list(deck.slides[0].shapes)
    names = [shape.name for shape in shapes]
    assert names == ["band", "logo plate", "logo"]
    plate, logo = shapes[1], shapes[2]
    assert plate.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE
    assert logo.shape_type == MSO_SHAPE_TYPE.PICTURE
    # The logo sits inside its panel.
    assert plate.left <= logo.left and plate.top <= logo.top
    assert logo.left + logo.width <= plate.left + plate.width
    assert logo.top + logo.height <= plate.top + plate.height
