"""A poster as a PowerPoint file whose layers stay editable, for finishing in Canva.

The agency's team already works in Canva. Handing them a flat PNG would throw away
the one thing Stratum has that an image model does not: every word, shape and logo
is a separate object. Canva imports .pptx with each object intact, so this writes
one slide per poster where each layer of the spec becomes one PowerPoint object,
in the same stacking order and named after the layer:

    image    -> a picture, cropped exactly as the poster crops it ("cover"), or
                fitted inside its box ("contain"); a blur is baked into the pixels,
                a dim is a separate black rectangle so it can be deleted
    cutout   -> a picture, never cropped, transparency kept
    logo     -> the client's original file, never cropped, never recoloured
    shape    -> a native rectangle, rounded rectangle, oval, pill or 16-point star
                in the brand colour; a faded band is a picture of that colour
                fading out, since PowerPoint has no fade that Canva imports
    text     -> a text box at the size and line breaks the renderer chose, so the
                file opens looking like the approved poster and stays editable,
                with its outline or shadow as PowerPoint's own

A photo the poster fades, mirrors or changes the light of (M10i) is baked: exactly
the part shown, at its box, with those changes in the pixels (and its darkening, so
the fade covers both). The original photo is untouched in Stratum.

Geometry and text fitting come from render/layout.py, the code the renderer uses.
Pixels map to EMU at 96 dpi and to points at 0.75 pt per pixel, which is how Canva
turns a PowerPoint page back into pixels.
"""

from __future__ import annotations

import io
from collections.abc import Mapping
from pathlib import Path
from typing import Any

import numpy as np
from PIL import Image, ImageFilter, ImageOps
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Emu, Pt
from stratum_spec import (
    BrandKit,
    CutoutLayer,
    ImageLayer,
    LogoLayer,
    PosterSpec,
    ShapeLayer,
    TextLayer,
)

from stratum_worker.render.layout import box_px, contained, crop_window, place_text, plated
from stratum_worker.render.svg import (
    BLUR_SCALE,
    DIM_FILL,
    PLACEHOLDER_FILL,
    WARM_SHIFT,
    effect_colour,
)

EMU_PER_PX = 9525  # 914400 EMU per inch / 96 px per inch
PT_PER_PX = 0.75  # 72 pt per inch / 96 px per inch
# Photos are kept at up to twice the size they are shown: sharp on a retina screen,
# without a 40 MB file for a 1080 px post.
MAX_SCALE = 2.0
JPEG_QUALITY = 90
ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
SHAPES = {
    "rect": MSO_SHAPE.RECTANGLE,
    "rounded": MSO_SHAPE.ROUNDED_RECTANGLE,
    "ellipse": MSO_SHAPE.OVAL,
    "pill": MSO_SHAPE.ROUNDED_RECTANGLE,
    "burst": MSO_SHAPE.STAR_16_POINT,
}
AnyImage = ImageLayer | CutoutLayer | LogoLayer
Geom = tuple[float, float, float, float]


class ExportError(ValueError):
    """The spec cannot be written as a PowerPoint file; the message says why."""


def _emu(px: float) -> Emu:
    return Emu(round(px * EMU_PER_PX))


def _rgb(hex_colour: str) -> RGBColor:
    digits = hex_colour.lstrip("#")
    red, green, blue = (int(digits[i : i + 2], 16) for i in (0, 2, 4))
    # python-pptx 1.0 ships annotations, but not on RGBColor's constructor.
    return RGBColor(red, green, blue)  # type: ignore[no-untyped-call]


def _child(parent: Any, tag: str, **attrs: str) -> None:
    element = OxmlElement(tag)
    for name, value in attrs.items():
        element.set(name, value)
    parent.append(element)


def _alpha(element: Any, opacity: float) -> None:
    """Make the first solid colour under `element` translucent. 100000 is opaque."""
    if opacity >= 1.0:
        return
    colour = element.find(f".//{qn('a:srgbClr')}")
    if colour is not None:
        _child(colour, "a:alpha", val=str(round(opacity * 100000)))


def _picture_alpha(picture: Any, opacity: float) -> None:
    if opacity >= 1.0:
        return
    blip = picture._element.find(f".//{qn('a:blip')}")
    if blip is not None:
        _child(blip, "a:alphaModFix", amt=str(round(opacity * 100000)))


def _plain(shape: Any) -> None:
    """No outline, no shadow, no theme styling: the shape is exactly its fill.

    A new PowerPoint shape carries a <p:style> pointing at the theme's line, fill and
    shadow. Importers apply it differently (LibreOffice drew a drop shadow under the
    brand rule), so it is removed and every property the poster has is set directly.
    """
    shape.line.fill.background()
    style = shape._element.find(qn("p:style"))
    if style is not None:
        shape._element.remove(style)


def _encode(image: Image.Image) -> bytes:
    """PNG when there is transparency to keep, JPEG otherwise (about a tenth the size)."""
    buffer = io.BytesIO()
    transparent = image.mode in ("RGBA", "LA") or (
        image.mode == "P" and "transparency" in image.info
    )
    if transparent:
        image.convert("RGBA").save(buffer, "PNG", optimize=True)
    else:
        image.convert("RGB").save(buffer, "JPEG", quality=JPEG_QUALITY, optimize=True)
    return buffer.getvalue()


def _limit(image: Image.Image, shown_w: float, shown_h: float) -> Image.Image:
    """Shrink an image that is far larger than it will ever be shown."""
    ratio = min(1.0, shown_w * MAX_SCALE / image.width, shown_h * MAX_SCALE / image.height)
    if ratio >= 1.0:
        return image
    size = (max(1, round(image.width * ratio)), max(1, round(image.height * ratio)))
    return image.resize(size, Image.Resampling.LANCZOS)


def _cover(slide: Any, layer: ImageLayer, spec: PosterSpec, image: Image.Image, geom: Geom) -> Any:
    """Fill the box, trimming the photo as the poster does: its crop, or evenly."""
    x, y, w, h = geom
    window = crop_window(image.size, (w, h), layer.crop) if layer.crop is not None else None
    scale = max(w / image.width, h / image.height) if window is None else w / window[2]
    if layer.blur > 0:
        # The renderer blurs at canvas scale; the same blur measured in source pixels.
        sigma = layer.blur * BLUR_SCALE * min(spec.canvas.size) / scale
        image = image.convert("RGB").filter(ImageFilter.GaussianBlur(sigma))
    source_size = image.size
    image = _limit(image, image.width * scale, image.height * scale)
    picture = slide.shapes.add_picture(
        io.BytesIO(_encode(image)), _emu(x), _emu(y), _emu(w), _emu(h)
    )
    if window is not None:
        # The crop window, as the share of the photo trimmed from each side. Shares,
        # not pixels, so the shrink `_limit` may have applied does not matter.
        (win_x, win_y, win_w, win_h), (full_w, full_h) = window, source_size
        picture.crop_left = win_x / full_w
        picture.crop_right = 1 - (win_x + win_w) / full_w
        picture.crop_top = win_y / full_h
        picture.crop_bottom = 1 - (win_y + win_h) / full_h
        return picture
    image_ratio, box_ratio = image.width / image.height, w / h
    if image_ratio > box_ratio:  # wider than the box: trim the sides
        side = (1 - box_ratio / image_ratio) / 2
        picture.crop_left = picture.crop_right = side
    elif image_ratio < box_ratio:  # taller than the box: trim top and bottom
        side = (1 - image_ratio / box_ratio) / 2
        picture.crop_top = picture.crop_bottom = side
    return picture


def _ramp(length: int) -> np.ndarray:
    """1 to 0 over `length` steps, eased as the renderer eases a fade (smoothstep)."""
    t = np.linspace(0.0, 1.0, max(length, 1))
    out: np.ndarray = 1.0 - t * t * (3.0 - 2.0 * t)
    return out


def _fade_alpha(size: tuple[int, int], fades: Any) -> np.ndarray:
    """The shown share of every pixel of a faded box (1 = all there)."""
    width, height = size
    alpha = np.ones((height, width))
    for fade in fades:
        span = round(fade.length * (height if fade.edge in ("top", "bottom") else width))
        ramp = _ramp(span)  # from shown (inside) to gone (at the edge)
        if fade.edge == "bottom":
            alpha[height - span :, :] *= ramp[:, None]
        elif fade.edge == "top":
            alpha[:span, :] *= ramp[::-1][:, None]
        elif fade.edge == "right":
            alpha[:, width - span :] *= ramp[None, :]
        else:
            alpha[:, :span] *= ramp[::-1][None, :]
    return alpha


def _baked(layer: ImageLayer, spec: PosterSpec, image: Image.Image, geom: Geom) -> Image.Image:
    """Exactly what the poster shows in this photo's box, with its fades, mirroring,
    light and colour, blur and darkening in the pixels."""
    _, _, w, h = geom
    shape: Image.Image | None = None  # where a "contain" photo is, in its box
    if layer.fit == "cover":
        win_x, win_y, win_w, win_h = crop_window(image.size, (w, h), layer.crop)
        image = image.crop((round(win_x), round(win_y), round(win_x + win_w), round(win_y + win_h)))
    else:  # the whole photo, centred in the box, as "contain" draws it
        ratio = min(w / image.width, h / image.height)
        size = (max(1, round(image.width * ratio)), max(1, round(image.height * ratio)))
        canvas = Image.new("RGB", (max(1, round(w)), max(1, round(h))))
        shape = Image.new("L", canvas.size, 0)
        at = ((canvas.width - size[0]) // 2, (canvas.height - size[1]) // 2)
        canvas.paste(image.convert("RGB").resize(size), at)
        shape.paste(255, (*at, at[0] + size[0], at[1] + size[1]))
        image = canvas
    scale = min(MAX_SCALE, max(1.0, image.width / w))
    out_w, out_h = max(1, round(w * scale)), max(1, round(h * scale))
    shown = image.convert("RGB").resize((out_w, out_h), Image.Resampling.LANCZOS)
    if shape is not None:
        shape = shape.resize((out_w, out_h))
    if layer.flip:
        shown = ImageOps.mirror(shown)
        shape = ImageOps.mirror(shape) if shape is not None else None
    if layer.blur > 0:
        sigma = layer.blur * BLUR_SCALE * min(spec.canvas.size) * scale
        shown = shown.filter(ImageFilter.GaussianBlur(sigma))
    pixels = np.asarray(shown, dtype=np.float64) / 255.0
    if layer.adjust is not None:
        adjust = layer.adjust
        grey = pixels @ np.array([0.2126, 0.7152, 0.0722])  # the SVG saturate weights
        pixels = grey[..., None] + (pixels - grey[..., None]) * (1.0 + adjust.saturation)
        slope = 1.0 + adjust.contrast
        intercept = 0.5 * (1.0 - slope) + adjust.brightness
        warm = WARM_SHIFT * adjust.warmth
        slopes = np.array([slope * (1 + warm), slope, slope * (1 - warm)])
        pixels = pixels * slopes + intercept
    if layer.dim > 0:
        pixels = pixels * (1.0 - layer.dim)
    rgb = np.clip(pixels * 255.0, 0, 255).astype(np.uint8)
    baked = Image.fromarray(rgb, "RGB")
    if layer.fades or shape is not None:
        alpha = _fade_alpha((out_w, out_h), layer.fades)
        if shape is not None:
            alpha = alpha * (np.asarray(shape, dtype=np.float64) / 255.0)
        baked.putalpha(Image.fromarray(np.clip(alpha * 255.0, 0, 255).astype(np.uint8), "L"))
    return baked


def _faded_band(slide: Any, layer: ShapeLayer, kit: BrandKit, geom: Geom) -> Any:
    """A band of colour that fades out, as a picture: PowerPoint's gradient fills do
    not survive Canva's import."""
    x, y, w, h = geom
    width, height = max(1, round(w)), max(1, round(h))
    colour = Image.new("RGB", (width, height), kit.color(layer.fill))
    alpha = _fade_alpha((width, height), layer.fades) * layer.opacity * 255.0
    colour.putalpha(Image.fromarray(np.clip(alpha, 0, 255).astype(np.uint8), "L"))
    picture = slide.shapes.add_picture(
        io.BytesIO(_encode(colour)), _emu(x), _emu(y), _emu(w), _emu(h)
    )
    picture.name = layer.id
    return picture


def _contain(slide: Any, image: Image.Image, geom: Geom) -> Any:
    """The whole image, as large as fits, centred in the box. Logos and cutouts always."""
    x, y, shown_w, shown_h = contained(geom, image.size)
    image = _limit(image, shown_w, shown_h)
    return slide.shapes.add_picture(
        io.BytesIO(_encode(image)), _emu(x), _emu(y), _emu(shown_w), _emu(shown_h)
    )


def _rectangle(slide: Any, geom: Geom, colour: str, opacity: float = 1.0) -> Any:
    x, y, w, h = geom
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, _emu(x), _emu(y), _emu(w), _emu(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(colour)
    _alpha(shape._element.spPr, opacity)
    _plain(shape)
    return shape


def _plate(slide: Any, layer: LogoLayer, kit: BrandKit, panel: Geom, radius: float) -> Any:
    """The panel behind a logo, as its own shape so a designer can restyle or delete it."""
    assert layer.plate is not None
    x, y, w, h = panel
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, _emu(x), _emu(y), _emu(w), _emu(h))
    shape.adjustments[0] = min(0.5, radius / min(w, h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(kit.color(layer.plate.fill))
    _alpha(shape._element.spPr, layer.plate.opacity * layer.opacity)
    _plain(shape)
    shape.name = f"{layer.id} plate"
    return shape


def _image_layer(
    slide: Any,
    layer: AnyImage,
    spec: PosterSpec,
    kit: BrandKit,
    assets: Mapping[str, str],
    geom: Geom,
) -> list[Any]:
    path = assets.get(layer.asset) if layer.asset else None
    if path is None:
        # An empty photo slot stays visible as a grey box, as on the poster.
        placeholder = _rectangle(slide, geom, PLACEHOLDER_FILL)
        placeholder.name = f"{layer.id} (empty {layer.type})"
        return [placeholder]
    with Image.open(path) as source:
        source.load()
        image = source.copy()
    made: list[Any] = []
    if isinstance(layer, LogoLayer) and layer.plate is not None:
        geom, panel, radius = plated(geom, image.size, layer.plate)
        made.append(_plate(slide, layer, kit, panel, radius))
    baked = isinstance(layer, ImageLayer) and (
        layer.fades or layer.flip or layer.adjust is not None
    )
    if isinstance(layer, ImageLayer) and baked:
        x, y, w, h = geom
        data = io.BytesIO(_encode(_baked(layer, spec, image, geom)))
        picture = slide.shapes.add_picture(data, _emu(x), _emu(y), _emu(w), _emu(h))
    elif isinstance(layer, ImageLayer) and layer.fit == "cover":
        picture = _cover(slide, layer, spec, image, geom)
    else:
        picture = _contain(slide, image, geom)
    picture.name = layer.id
    _picture_alpha(picture, layer.opacity)
    made.append(picture)
    if isinstance(layer, ImageLayer) and layer.dim > 0 and not baked:
        dim = _rectangle(slide, geom, DIM_FILL, layer.dim * layer.opacity)
        dim.name = f"{layer.id} dim"
        made.append(dim)
    return made


def _shape_layer(slide: Any, layer: ShapeLayer, kit: BrandKit, geom: Geom) -> Any:
    if layer.fades:
        return _faded_band(slide, layer, kit, geom)
    x, y, w, h = geom
    shape = slide.shapes.add_shape(SHAPES[layer.shape], _emu(x), _emu(y), _emu(w), _emu(h))
    if layer.shape == "rounded":
        # Both formats measure the corner as a fraction of the shorter side.
        shape.adjustments[0] = min(0.5, layer.radius)
    elif layer.shape == "pill":
        shape.adjustments[0] = 0.5
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(kit.color(layer.fill))
    _alpha(shape._element.spPr, layer.opacity)
    _plain(shape)
    shape.name = layer.id
    return shape


def _in_front(slide: Any, text: TextLayer, spec: PosterSpec, assets: Mapping[str, str]) -> Any:
    """Depth (schema 6): the photo's subject, cut out, placed exactly over its own photo
    and above these words, so the words sit behind it as on the poster.

    A text box cannot be masked in PowerPoint, so instead of hiding letters the subject
    is drawn again on top: the photo's pixels with the subject matte as transparency,
    dimmed and blurred as the photo is, trimmed by the same crop. A designer can move
    or delete it like any picture. None when the words are simply in front.
    """
    if text.behind is None:
        return None
    photo = next((one for one in spec.layers if one.id == text.behind), None)
    if not isinstance(photo, ImageLayer) or not photo.visible or photo.fit != "cover":
        return None
    picture = assets.get(photo.asset) if photo.asset else None
    matte = assets.get(photo.subject) if photo.subject else None
    if picture is None or matte is None:
        return None
    with Image.open(picture) as source:
        image = source.convert("RGB")
    with Image.open(matte) as source:
        alpha = source.convert("L").resize(image.size, Image.Resampling.BILINEAR)
    geom = box_px(spec, photo)
    _, _, w, h = geom
    if photo.blur > 0:
        window = crop_window(image.size, (w, h), photo.crop) if photo.crop is not None else None
        scale = max(w / image.width, h / image.height) if window is None else w / window[2]
        sigma = photo.blur * BLUR_SCALE * min(spec.canvas.size) / scale
        image = image.filter(ImageFilter.GaussianBlur(sigma))
    if photo.dim > 0:
        image = Image.eval(image, lambda value: round(value * (1.0 - photo.dim)))
    if photo.opacity < 1.0:
        alpha = Image.eval(alpha, lambda value: round(value * photo.opacity))
    image.putalpha(alpha)
    front = _cover(slide, photo.model_copy(update={"blur": 0.0}), spec, image, geom)
    front.name = f"{text.id} behind: subject of {photo.id}"
    return front


def _text_layer(slide: Any, layer: TextLayer, spec: PosterSpec, kit: BrandKit, root: Path) -> Any:
    placed = place_text(layer, spec, kit, root)
    x, y, w, h = placed.box
    box = slide.shapes.add_textbox(_emu(x), _emu(y), _emu(w), _emu(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = Emu(0)
    paragraph = frame.paragraphs[0]
    paragraph.alignment = ALIGN[placed.align]
    paragraph.line_spacing = Pt(round(placed.line_height_px * PT_PER_PX, 2))
    # One paragraph, with the renderer's own line breaks: it opens exactly as the
    # poster was drawn, and a designer editing it gets normal reflow from there.
    for index, line in enumerate(placed.lines):
        if index:
            paragraph.add_line_break()
        run = paragraph.add_run()
        run.text = line
        font = run.font
        font.name = placed.family
        font.size = Pt(round(placed.size_px * PT_PER_PX, 2))  # PowerPoint keeps 1/100 pt
        font.bold = placed.weight >= 600
        font.color.rgb = _rgb(placed.color)
        _alpha(run._r, layer.opacity)
        _text_effect(run, layer, kit, placed.size_px)
    box.name = layer.id
    return box


def _text_effect(run: Any, layer: TextLayer, kit: BrandKit, size_px: float) -> None:
    """The words' outline or shadow (M10i) as PowerPoint's own, so it stays editable."""
    effect = layer.effect
    if effect is None:
        return
    colour = effect_colour(layer, kit).lstrip("#").upper()
    properties = run._r.get_or_add_rPr()
    if effect.kind == "outline":
        width = round(size_px * 0.1 * effect.strength * EMU_PER_PX)
        line = OxmlElement("a:ln")
        line.set("w", str(width))
        fill = OxmlElement("a:solidFill")
        rgb = OxmlElement("a:srgbClr")
        rgb.set("val", colour)
        fill.append(rgb)
        line.append(fill)
        properties.insert(0, line)  # an outline comes before the fill in a run
        return
    shadow_list = OxmlElement("a:effectLst")
    shadow = OxmlElement("a:outerShdw")
    shadow.set("blurRad", str(round(size_px * 0.16 * effect.strength * EMU_PER_PX)))
    shadow.set("dist", str(round(size_px * 0.04 * EMU_PER_PX)))
    shadow.set("dir", "5400000")  # straight down
    shadow.set("algn", "t")
    shadow.set("rotWithShape", "0")
    rgb = OxmlElement("a:srgbClr")
    rgb.set("val", colour)
    alpha = OxmlElement("a:alpha")
    alpha.set("val", str(round((0.35 + 0.5 * effect.strength) * 100000)))
    rgb.append(alpha)
    shadow.append(rgb)
    shadow_list.append(shadow)
    fill = properties.find(qn("a:solidFill"))
    if fill is not None:
        fill.addnext(shadow_list)
    else:
        properties.append(shadow_list)


def export_pptx(
    spec: PosterSpec,
    kit: BrandKit,
    *,
    assets: Mapping[str, str] | None = None,
    root: Path | None = None,
    title: str = "",
    notes: str = "",
) -> bytes:
    """One poster as a one-slide .pptx, a PowerPoint object per visible layer.

    `assets` maps a SHA-256 to a local file, as for the renderer. `notes` becomes
    the slide's speaker notes, where Canva shows it beside the design.
    """
    assets = assets or {}
    root = root or Path.cwd()
    width, height = spec.canvas.size
    deck = Presentation()
    deck.slide_width, deck.slide_height = _emu(width), _emu(height)
    deck.core_properties.title = title[:200]
    deck.core_properties.author = "Stratum"
    slide = deck.slides.add_slide(deck.slide_layouts[6])  # the blank layout
    for layer in spec.layers:
        if not layer.visible:
            continue
        geom = box_px(spec, layer)
        made: list[Any]
        if isinstance(layer, TextLayer):
            made = [_text_layer(slide, layer, spec, kit, root)]
        elif isinstance(layer, ShapeLayer):
            made = [_shape_layer(slide, layer, kit, geom)]
        elif isinstance(layer, ImageLayer | CutoutLayer | LogoLayer):
            made = _image_layer(slide, layer, spec, kit, assets, geom)
        else:  # pragma: no cover - the spec's union is closed
            raise ExportError(f"no PowerPoint mapping for layer type {layer.type!r}")
        for shape in made:
            shape.rotation = layer.box.rotation
        if isinstance(layer, TextLayer):
            _in_front(slide, layer, spec, assets)
    if notes:
        slide.notes_slide.notes_text_frame.text = notes
    buffer = io.BytesIO()
    deck.save(buffer)
    return buffer.getvalue()
